"""GUI input modes and ECU pairing; skip in installations without optional Qt."""

from __future__ import annotations

import os
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

try:
    from PySide6.QtCore import QSettings, QTimer
    from PySide6.QtTest import QTest
    from PySide6.QtWidgets import QApplication, QDialog
except ImportError:
    HAS_QT = False
else:
    HAS_QT = True
    from dbc_compare_tool.ui.main_window import (
        CompareWorker,
        EcuDiscoveryWorker,
        EcuPairingDialog,
        MainWindow,
        ManualPairingDialog,
    )


_UI = "dbc_compare_tool.ui.main_window"
_DBC = '''VERSION ""
BU_: {sender} Receiver Spare
BO_ 100 Status: 8 {sender}
 SG_ Counter : 0|8@1+ (1,0) [0|255] "" Receiver
'''


@unittest.skipUnless(HAS_QT, "PySide6 is unavailable")
class GuiInputTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.old = self.root / "old.dbc"
        self.new = self.root / "new.dbc"
        self.old.write_text(_DBC.format(sender="OldNode"), encoding="utf-8")
        self.new.write_text(_DBC.format(sender="NewNode"), encoding="utf-8")
        self.settings = QSettings(str(self.root / "settings.ini"), QSettings.Format.IniFormat)
        with patch(f"{_UI}.QSettings", return_value=self.settings):
            self.window = MainWindow()
        self.addCleanup(self._close_window)

    def _close_window(self):
        for worker in (self.window.worker, self.window.export_worker, self.window.ecu_worker):
            if worker:
                worker.wait(5000)
        self.window.close()
        self.window.deleteLater()
        self.app.processEvents()

    def _files(self):
        self.window.input_mode.setCurrentIndex(self.window.input_mode.findData("files"))
        self.window.old_input.setText(str(self.old))
        self.window.new_input.setText(str(self.new))

    def _save_maps(self):
        self.window._saved_pair_map = {"old.dbc": "new.dbc"}
        self.window._saved_pair_folders = self.window._current_folder_inputs()
        self.window._saved_ecu_map = {"OldNode": "NewNode"}
        self.window._saved_ecu_inputs = self.window._current_inputs()

    def _wait_until(self, condition):
        deadline = time.monotonic() + 10
        while not condition() and time.monotonic() < deadline:
            QTest.qWait(10)
        self.assertTrue(condition(), "Qt worker did not complete within 10 seconds")

    def test_file_browse_selects_dbc_and_disables_manual_file_pairing(self):
        self._files()
        with patch(f"{_UI}.QFileDialog.getOpenFileName", return_value=(str(self.old), "")) as pick:
            self.window._browse_buttons[0].click()
        self.assertIn("*.dbc", pick.call_args.args[3])
        self.assertEqual(self.window.old_input.text(), str(self.old))
        self.assertFalse(self.window.manual_pair_button.isEnabled())
        self.assertTrue(self.window.ecu_pair_button.isEnabled())
        self.assertEqual(self.window._validated_baselines(), (self.old, self.new))
        self.assertTrue(self.window.old_input._accepts_path(self.old))
        self.assertFalse(self.window.old_input._accepts_path(self.root))
        self.assertIn("DBC", self.window.old_label.text())

    def test_folder_browse_and_complete_manual_pairing_remain_available(self):
        folder = self.root / "baseline"
        folder.mkdir()
        (folder / "bus.dbc").write_text(_DBC.format(sender="OldNode"), encoding="utf-8")
        with patch(f"{_UI}.QFileDialog.getExistingDirectory", return_value=str(folder)):
            self.window._browse_buttons[0].click()
            self.window._browse_buttons[1].click()
        self.assertTrue(self.window.manual_pair_button.isEnabled())
        self.assertEqual(self.window._validated_baselines(), (folder, folder))
        self.assertTrue(self.window.old_input._accepts_path(folder))
        self.assertFalse(self.window.old_input._accepts_path(self.old))
        with patch(f"{_UI}.ManualPairingDialog") as dialog_type:
            dialog_type.return_value.exec.return_value = QDialog.DialogCode.Accepted
            dialog_type.return_value.pair_map.return_value = {"bus.dbc": "bus.dbc"}
            self.assertTrue(self.window._open_manual_pairing())
        self.assertEqual(self.window._complete_manual_pairing(), {"bus.dbc": "bus.dbc"})
        self.window._saved_pair_map = {"bus.dbc": None}
        self.assertEqual(self.window._complete_manual_pairing(), {"bus.dbc": None})

    def test_paths_and_input_mode_invalidate_both_saved_maps(self):
        self._files()
        self._save_maps()
        self.window.old_input.setText(str(self.new))
        self.assertIsNone(self.window._saved_pair_map)
        self.assertEqual(self.window._saved_ecu_map, {})
        self._save_maps()
        self.window.input_mode.setCurrentIndex(0)
        self.assertIsNone(self.window._saved_pair_map)
        self.assertIsNone(self.window._saved_ecu_inputs)
        self.assertEqual(self.window._saved_ecu_map, {})

    def test_file_mode_rejects_folder_and_non_dbc_inputs(self):
        self._files()
        for invalid in (self.root, self.root / "missing.dbc", self.root / "notes.txt"):
            with self.subTest(invalid=invalid):
                self.window.new_input.setText(str(invalid))
                with patch(f"{_UI}.QMessageBox.warning") as warning:
                    self.assertIsNone(self.window._validated_baselines())
                warning.assert_called_once()

    def test_input_mode_is_saved_and_restored_with_explicit_output(self):
        self._files()
        self.window.output_input.setText(str(self.root / "custom.xlsx"))
        self.window._save_paths()
        with patch(f"{_UI}.QSettings", return_value=self.settings):
            restored = MainWindow()
        try:
            self.assertTrue(restored._files_mode())
            self.assertEqual(restored.new_input.text(), str(self.new))
            self.assertEqual(restored.output_input.text(), str(self.root / "custom.xlsx"))
            self.assertFalse(restored.manual_pair_button.isEnabled())
        finally:
            restored.close()
            restored.deleteLater()

    def test_compare_files_with_ecu_map_exports_report_through_both_workers(self):
        self._files()
        self._save_maps()
        self.window.chk_include_unchanged.setChecked(True)
        report = self.root / "report.xlsx"
        self.window.output_input.setText(str(report))
        results = []
        with patch(f"{_UI}.QMessageBox.critical") as error:
            self.window._run_compare()
            self.window.worker.compared.connect(results.append)
            self.assertFalse(self.window.input_mode.isEnabled())
            self.assertFalse(self.window.ecu_pair_button.isEnabled())
            self.assertFalse(self.window.new_input.isEnabled())
            self.assertTrue(all(not button.isEnabled() for button in self.window._browse_buttons))
            self._wait_until(lambda: self.window.run_button.isEnabled())
        error.assert_not_called()
        self.assertEqual(self.window.last_report, report)
        self.assertTrue(report.is_file())
        self.assertEqual(results[0].ecu_map, {"OldNode": "NewNode"})
        self.assertEqual(results[0].summary()["Total Changes"], 0)
        self.assertTrue(results[0].include_unchanged)
        self.assertIsNone(self.window.worker.pair_map)
        self.assertTrue(self.window.input_mode.isEnabled())
        self.assertFalse(self.window.manual_pair_button.isEnabled())
        self.assertTrue(self.window.progress.isHidden())

    def test_gui_export_splits_names_and_omits_added_removed_arrows(self):
        old_folder, new_folder = self.root / "old", self.root / "new"
        old_folder.mkdir()
        new_folder.mkdir()
        old_dbc = _DBC.format(sender="Controller") + ' SG_ Removed : 8|8@1+ (1,0) [0|255] "" Receiver\n'
        new_dbc = _DBC.format(sender="Controller").replace("Status:", "StatusNew:").replace("Counter :", "CounterNew :")
        new_dbc += ' SG_ Added : 32|32@0- (0.01,-100) [-100|100] "degC" Spare\n'
        (old_folder / "bus.dbc").write_text(old_dbc, encoding="utf-8")
        (new_folder / "bus.dbc").write_text(new_dbc, encoding="utf-8")
        self.window.old_input.setText(str(old_folder))
        self.window.new_input.setText(str(new_folder))
        report = self.root / "hotfix.xlsx"
        self.window.output_input.setText(str(report))
        with patch(f"{_UI}.QMessageBox.critical") as error:
            self.window.run_button.click()
            self._wait_until(lambda: self.window.run_button.isEnabled())
        error.assert_not_called()
        self.assertEqual(self.window.last_report, report)
        workbook = load_workbook(report)
        self.addCleanup(workbook.close)
        kinds = set()
        for sheet_name, entity in (("Message Details", "Message"), ("Signal Details", "Signal")):
            sheet = workbook[sheet_name]
            headers = next(sheet.iter_rows(values_only=True))
            self.assertIn(f"Old {entity} Name", headers)
            self.assertIn(f"New {entity} Name", headers)
            for cells in sheet.iter_rows(min_row=2, values_only=True):
                row = dict(zip(headers, cells))
                kind = row["Change Type"]
                kinds.add(kind)
                if kind in {"Added", "Removed"}:
                    self.assertTrue(all("->" not in str(value) for value in row.values()))
                elif kind == "Renamed":
                    self.assertTrue(row[f"Old {entity} Name"])
                    self.assertTrue(row[f"New {entity} Name"])
                    self.assertNotEqual(row[f"Old {entity} Name"], row[f"New {entity} Name"])
        self.assertTrue({"Added", "Removed", "Renamed"}.issubset(kinds))

    def test_ecu_discovery_runs_in_worker_and_saves_selected_mapping(self):
        self._files()
        with patch(f"{_UI}.EcuPairingDialog") as dialog_type:
            dialog_type.return_value.exec.return_value = QDialog.DialogCode.Accepted
            dialog_type.return_value.pair_map.return_value = {"OldNode": "NewNode"}
            self.window._open_ecu_pairing()
            self.assertIsInstance(self.window.ecu_worker, EcuDiscoveryWorker)
            self.assertFalse(self.window.run_button.isEnabled())
            self._wait_until(lambda: self.window.run_button.isEnabled())
            self.assertEqual(dialog_type.call_args.args[1], ["OldNode", "Receiver", "Spare"])
            self.assertEqual(dialog_type.call_args.args[2], ["NewNode", "Receiver", "Spare"])
        self.assertEqual(self.window._saved_ecu_map, {"OldNode": "NewNode"})
        self.assertEqual(self.window._saved_ecu_inputs, self.window._current_inputs())

    def test_ecu_pairing_stops_loading_before_modal_editor_opens(self):
        original_exec = EcuPairingDialog.exec
        old_folder, new_folder = self.root / "old", self.root / "new"
        old_folder.mkdir()
        new_folder.mkdir()
        (old_folder / "bus.dbc").write_text(self.old.read_text(), encoding="utf-8")
        (new_folder / "bus.dbc").write_text(self.new.read_text(), encoding="utf-8")
        for mode in ("files", "folders"):
            for save in (False, True):
                with self.subTest(mode=mode, save=save):
                    self._files()
                    if mode == "folders":
                        self.window.input_mode.setCurrentIndex(0)
                        self.window.old_input.setText(str(old_folder))
                        self.window.new_input.setText(str(new_folder))
                    self.window.last_report = self.root / "previous.xlsx"
                    self._save_maps()
                    self.window._saved_ecu_map = {}
                    observed = []

                    def inspect_editor(dialog):
                        observed.append((
                            self.window.run_status.text(), self.window.progress.maximum(),
                            self.window.progress.isHidden(), self.window.run_button.isEnabled(),
                            self.window.old_input.isEnabled(),
                        ))
                        if save:
                            row = next(
                                row for row in range(dialog.table.rowCount())
                                if dialog.table.item(row, 0).text() == "OldNode"
                            )
                            combo = dialog.table.cellWidget(row, 1)
                            combo.setCurrentIndex(combo.findData("NewNode"))
                        QTimer.singleShot(0, dialog._on_save if save else dialog.reject)
                        return original_exec(dialog)

                    with patch.object(EcuPairingDialog, "exec", inspect_editor):
                        self.window._open_ecu_pairing()
                        self.assertEqual(self.window.run_status.text(), "Loading ECU nodes…")
                        self.assertEqual(self.window.progress.maximum(), 0)
                        self.assertFalse(self.window.progress.isHidden())
                        self._wait_until(lambda: self.window.run_button.isEnabled())
                    self.assertEqual(observed, [("Editing ECU pairing…", 1, True, False, False)])
                    self.assertEqual(self.window.run_status.text(), "Ready to compare")
                    self.assertTrue(self.window.progress.isHidden())
                    self.assertTrue(self.window.open_button.isEnabled())
                    self.assertEqual(self.window.manual_pair_button.isEnabled(), mode == "folders")
                    self.assertEqual(
                        self.window._saved_ecu_map,
                        {"OldNode": "NewNode"} if save else {},
                    )

    def test_bulk_rename_exports_correct_rows_through_gui(self):
        from tests.test_bulk_rename import NEW_DBC, OLD_DBC, SUFFIXES

        self.old.write_text(OLD_DBC, encoding="utf-8")
        self.new.write_text(NEW_DBC, encoding="utf-8")
        self._files()
        report = self.root / "bulk.xlsx"
        self.window.output_input.setText(str(report))
        with patch(f"{_UI}.QMessageBox.critical") as error:
            self.window._run_compare()
            self._wait_until(lambda: self.window.run_button.isEnabled())
        error.assert_not_called()
        workbook = load_workbook(report)
        self.addCleanup(workbook.close)
        values = list(workbook["Signal Details"].values)
        rows = [dict(zip(values[0], row)) for row in values[1:]]
        self.assertEqual(len(rows), len(SUFFIXES))
        self.assertTrue(all(row["DBC File"] == "old.dbc->new.dbc" and row["Change Type"] == "Renamed" for row in rows))
        self.assertNotIn("Changed Properties", workbook.sheetnames)
        self.assertTrue(all(row["Changed Properties"] is None for row in rows))

    def test_gui_exports_multiple_property_changes_in_each_detail_sheet(self):
        self._files()
        old = _DBC.format(sender="OldNode")
        self.old.write_text(old, encoding="utf-8")
        self.new.write_text(old.replace('0|8@1+ (1,0)', '1|8@1+ (2,0)')
                            + '\nCM_ BO_ 100 "Updated status";\n', encoding="utf-8")
        report = self.root / "properties.xlsx"
        self.window.output_input.setText(str(report))
        with patch(f"{_UI}.QMessageBox.critical") as error:
            self.window._run_compare()
            self._wait_until(lambda: self.window.run_button.isEnabled())
        error.assert_not_called()
        self.window.worker.wait(5000)
        self.window.export_worker.wait(5000)
        workbook = load_workbook(report)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames[:3], ["Summary", "Message Details", "Signal Details"])
        self.assertNotIn("DBC Overview", workbook.sheetnames)
        for name, expected in (
            ("Message Details", "Description: ->Updated status"),
            ("Signal Details", "Start Bit: 0->1; Factor: 1->2"),
        ):
            values = list(workbook[name].values)
            self.assertEqual(len(values), 2)
            row = dict(zip(values[0], values[1]))
            self.assertEqual(row["Changed Properties"], expected)

    def test_gui_exports_same_name_id_change_and_precise_values_without_overview(self):
        self._files()
        old = _DBC.format(sender="OldNode").replace("(1,0)", "(0.1234567,0)")
        self.old.write_text(old, encoding="utf-8")
        self.new.write_text(old.replace("BO_ 100 ", "BO_ 400 ").replace("0.1234567", "0.1234568"),
                            encoding="utf-8")
        report = self.root / "precise.xlsx"
        self.window.output_input.setText(str(report))
        with patch(f"{_UI}.QMessageBox.critical") as error:
            self.window._run_compare()
            self._wait_until(lambda: self.window.run_button.isEnabled())
        error.assert_not_called()
        workbook = load_workbook(report)
        self.addCleanup(workbook.close)
        for name, expected in (("Message Details", "CAN ID: 100->400"),
                               ("Signal Details", "Factor: 0.1234567->0.1234568")):
            values = list(workbook[name].values)
            self.assertEqual(len(values), 2)
            row = dict(zip(values[0], values[1]))
            self.assertEqual(row["Change Type"], "Modified")
            self.assertEqual(row["Changed Properties"], expected)
        self.assertEqual(workbook.sheetnames[:3], ["Summary", "Message Details", "Signal Details"])
        self.assertNotIn("DBC Overview", workbook.sheetnames)

    def test_gui_folder_report_keeps_overview_for_one_pair(self):
        old_folder, new_folder = self.root / "old", self.root / "new"
        for folder in (old_folder, new_folder):
            folder.mkdir()
            (folder / "bus.dbc").write_text(_DBC.format(sender="OldNode"), encoding="utf-8")
        self.window.input_mode.setCurrentIndex(self.window.input_mode.findData("folders"))
        self.window.old_input.setText(str(old_folder))
        self.window.new_input.setText(str(new_folder))
        report = self.root / "folder.xlsx"
        self.window.output_input.setText(str(report))
        with patch(f"{_UI}.QMessageBox.critical") as error:
            self.window._run_compare()
            self._wait_until(lambda: self.window.run_button.isEnabled())
        error.assert_not_called()
        workbook = load_workbook(report)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["Summary", "DBC Overview", "Message Details", "Signal Details"])
        self.assertEqual(workbook["DBC Overview"].max_row, 2)

    def test_gui_rejects_duplicate_message_names_without_exporting(self):
        self._files()
        old = _DBC.format(sender="OldNode")
        self.old.write_text(old + old.replace("BO_ 100 ", "BO_ 200 "), encoding="utf-8")
        report = self.root / "duplicate.xlsx"
        self.window.output_input.setText(str(report))
        with patch(f"{_UI}.QMessageBox.critical") as error:
            self.window._run_compare()
            self._wait_until(lambda: self.window.run_button.isEnabled())
        error.assert_called_once()
        self.assertIn("Duplicate message name 'Status'", error.call_args.args[-1])
        self.assertFalse(report.exists())

    def test_shifted_signal_pairing_confidence_is_exported_by_gui(self):
        from tests.test_signal_pairing import shifted_dbc

        old_folder, new_folder = self.root / "old", self.root / "new"
        old_folder.mkdir()
        new_folder.mkdir()
        for event_like in (False, True):
            old, new, pairs = shifted_dbc(event_like)
            self.old.write_text(old, encoding="utf-8")
            self.new.write_text(new, encoding="utf-8")
            (old_folder / "bus.dbc").write_text(old, encoding="utf-8")
            (new_folder / "bus.dbc").write_text(new, encoding="utf-8")
            for mode in ("files", "folders"):
                with self.subTest(event_like=event_like, mode=mode):
                    self._files()
                    if mode == "folders":
                        self.window.input_mode.setCurrentIndex(0)
                        self.window.old_input.setText(str(old_folder))
                        self.window.new_input.setText(str(new_folder))
                    report = self.root / "shifted.xlsx"
                    self.window.output_input.setText(str(report))
                    with patch(f"{_UI}.QMessageBox.critical") as error:
                        self.window._run_compare()
                        self._wait_until(lambda: self.window.run_button.isEnabled())
                    error.assert_not_called()
                    self.window.worker.wait(5000)
                    self.window.export_worker.wait(5000)
                    workbook = load_workbook(report)
                    try:
                        values = list(workbook["Signal Details"].values)
                        rows = [dict(zip(values[0], row)) for row in values[1:]]
                        self.assertEqual({(row["Old Signal Name"], row["New Signal Name"]) for row in rows},
                                         set(pairs.items()))
                        self.assertTrue(all(row["Change Type"] == "Renamed" for row in rows))
                        self.assertTrue(all(row["Confidence Level"] in {"Low", "Medium"} and row["Review Required"] == "Yes" for row in rows))
                        self.assertNotIn("Changed Properties", workbook.sheetnames)
                        for row in rows:
                            prop, value = row["Changed Properties"].split(": ")
                            self.assertEqual(prop, "Start Bit")
                            old_start, new_start = value.split("->")
                            self.assertEqual(int(new_start) - int(old_start), 1)
                    finally:
                        workbook.close()

    def test_ecu_parse_failure_restores_controls(self):
        self._files()
        self.new.write_text("invalid DBC", encoding="utf-8")
        with patch(f"{_UI}.QMessageBox.critical") as error:
            self.window._open_ecu_pairing()
            self._wait_until(lambda: self.window.run_button.isEnabled())
        error.assert_called_once()
        self.assertTrue(self.window.new_input.isEnabled())
        self.assertFalse(self.window.manual_pair_button.isEnabled())
        self.assertEqual(self.window.run_status.text(), "Ready to compare")
        self.assertEqual(self.window.progress.maximum(), 1)
        self.assertTrue(self.window.progress.isHidden())

    def test_stale_ecu_discovery_restores_idle_state_without_opening_editor(self):
        self._files()
        with patch(f"{_UI}.EcuDiscoveryWorker.start"):
            self.window._open_ecu_pairing()
        self.window.new_input.setText(str(self.old))
        with patch(f"{_UI}.EcuPairingDialog") as dialog_type:
            self.window._on_ecu_nodes_loaded(["OldNode"], ["NewNode"])
        dialog_type.assert_not_called()
        self.assertEqual(self.window.run_status.text(), "Ready to compare")
        self.assertTrue(self.window.progress.isHidden())
        self.assertTrue(self.window.run_button.isEnabled())

    def test_manual_pairing_editor_restores_previous_status_after_save_or_cancel(self):
        self.window.old_input.setText(str(self.root))
        self.window.new_input.setText(str(self.root))
        for result in (QDialog.DialogCode.Accepted, QDialog.DialogCode.Rejected):
            with self.subTest(result=result):
                self.window.run_status.setText("Report ready · 5 changes")

                def inspect_editor():
                    self.assertEqual(self.window.run_status.text(), "Editing DBC pairing…")
                    self.assertTrue(self.window.progress.isHidden())
                    return result

                with patch(f"{_UI}.ManualPairingDialog") as dialog_type:
                    dialog_type.return_value.exec.side_effect = inspect_editor
                    dialog_type.return_value.pair_map.return_value = {"old.dbc": "new.dbc"}
                    self.assertEqual(
                        self.window._open_manual_pairing(), result == QDialog.DialogCode.Accepted
                    )
                self.assertEqual(self.window.run_status.text(), "Report ready · 5 changes")

    def test_manual_folder_worker_keeps_ecu_map_and_original_names(self):
        old_folder, new_folder = self.root / "old", self.root / "new"
        old_folder.mkdir()
        new_folder.mkdir()
        (old_folder / "legacy.dbc").write_text(self.old.read_text(), encoding="utf-8")
        (new_folder / "current.dbc").write_text(self.new.read_text(), encoding="utf-8")
        worker = CompareWorker(
            old_folder, new_folder, {"legacy.dbc": "current.dbc"},
            include_unchanged=True, ecu_map={"OldNode": "NewNode"},
        )
        results, errors = [], []
        worker.compared.connect(results.append)
        worker.failed.connect(errors.append)
        worker.run()
        self.assertEqual(errors, [])
        self.assertEqual(results[0].summary()["Total Changes"], 0)
        self.assertEqual(results[0].message_changes[0].old_message.transmitter, "OldNode")
        self.assertEqual(results[0].message_changes[0].new_message.transmitter, "NewNode")

    def test_ecu_dialog_saves_one_or_two_pairs_with_other_nodes_kept(self):
        for mapping in ({"A": "B"}, {"A": "B", "C": "D"}):
            with self.subTest(mapping=mapping):
                nodes = ["A", "B", "C", "D", "E"]
                dialog = EcuPairingDialog(self.window, nodes, ["B", "C", "D", "E"])
                for row, node in enumerate(nodes):
                    if node in mapping:
                        combo = dialog.table.cellWidget(row, 1)
                        combo.setCurrentIndex(combo.findData(mapping[node]))
                with patch(f"{_UI}.QMessageBox.warning") as warning:
                    dialog._on_save()
                warning.assert_not_called()
                self.assertEqual(dialog.result(), QDialog.DialogCode.Accepted)
                self.assertEqual(dialog.pair_map(), mapping)
                restored = EcuPairingDialog(self.window, nodes, ["B", "C", "D", "E"], previous=dialog.pair_map())
                for row, node in enumerate(nodes):
                    self.assertEqual(restored.table.cellWidget(row, 1).currentData(), mapping.get(node))
                restored.deleteLater()
                dialog.deleteLater()

    def test_ecu_dialog_rejects_duplicate_selected_rename_targets(self):
        for target in ("C", "D"):
            with self.subTest(target=target):
                dialog = EcuPairingDialog(self.window, ["A", "B"], ["C", "D"])
                first = dialog.table.cellWidget(0, 1)
                first.setCurrentIndex(first.findData(target))
                second = dialog.table.cellWidget(1, 1)
                second.setCurrentIndex(second.findData(target))
                with patch(f"{_UI}.QMessageBox.warning") as warning:
                    dialog._on_save()
                warning.assert_called_once()
                self.assertNotEqual(dialog.result(), QDialog.DialogCode.Accepted)
                dialog.deleteLater()

    def test_partial_ecu_pairing_saves_and_exports_through_gui(self):
        self.old.write_text(_DBC.format(sender="OldNode").replace("Receiver Spare", "Receiver Spare NewNode"), encoding="utf-8")
        self._files()
        self.window.chk_include_unchanged.setChecked(True)
        report = self.root / "partial.xlsx"
        self.window.output_input.setText(str(report))
        original_exec = EcuPairingDialog.exec

        def save_partial(dialog):
            row = next(row for row in range(dialog.table.rowCount()) if dialog.table.item(row, 0).text() == "OldNode")
            combo = dialog.table.cellWidget(row, 1)
            combo.setCurrentIndex(combo.findData("NewNode"))
            QTimer.singleShot(0, dialog._on_save)
            return original_exec(dialog)

        with patch.object(EcuPairingDialog, "exec", save_partial), patch(f"{_UI}.QMessageBox.warning") as warning:
            self.window._open_ecu_pairing()
            self._wait_until(lambda: self.window.run_button.isEnabled())
        warning.assert_not_called()
        self.assertEqual(self.window._saved_ecu_map, {"OldNode": "NewNode"})
        with patch(f"{_UI}.QMessageBox.critical") as error:
            self.window._run_compare()
            self._wait_until(lambda: self.window.run_button.isEnabled())
        error.assert_not_called()
        self.assertEqual(self.window.worker.ecu_map, {"OldNode": "NewNode"})
        self.assertEqual(self.window.last_report, report)
        self.assertTrue(report.is_file())
        self.assertIn("0 changes", self.window.run_status.text())
        with report.open("rb") as stream:
            workbook = load_workbook(stream)
            try:
                for name in ("Message Details", "Signal Details"):
                    values = list(workbook[name].values)
                    row = dict(zip(values[0], values[1]))
                    self.assertEqual(row["ECU Node Tx"], "OldNode->NewNode")
                    self.assertEqual(row["ECU Node Rx"], "Receiver")
                    self.assertEqual(row["CAN ID"], "0x64")
                    self.assertEqual(row["Change Type"], "Unchanged")
            finally:
                workbook.close()

    def test_ecu_dialog_supports_swaps_and_clearing_saved_mapping(self):
        dialog = EcuPairingDialog(
            self.window, ["A", "B"], ["A", "B"], previous={"A": "B", "B": "A"}
        )
        dialog._on_save()
        self.assertEqual(dialog.result(), QDialog.DialogCode.Accepted)
        self.assertEqual(dialog.pair_map(), {"A": "B", "B": "A"})
        for row in range(2):
            dialog.table.cellWidget(row, 1).setCurrentIndex(0)
        dialog._on_save()
        self.assertEqual(dialog.pair_map(), {})
        dialog.deleteLater()

    def test_manual_dialog_still_rejects_duplicate_file_pairing(self):
        folder = self.root / "old"
        folder.mkdir()
        for name in ("one.dbc", "two.dbc"):
            (folder / name).write_text(self.old.read_text(), encoding="utf-8")
        dialog = ManualPairingDialog(self.window, folder, folder)
        dialog.table.cellWidget(0, 1).setCurrentIndex(1)
        dialog.table.cellWidget(1, 1).setCurrentIndex(1)
        with patch(f"{_UI}.QMessageBox.warning") as warning:
            dialog._on_save()
        warning.assert_called_once()
        self.assertNotEqual(dialog.result(), QDialog.DialogCode.Accepted)
        dialog.deleteLater()

    def test_same_corpus_matches_api_cli_and_gui_metadata_after_excel_roundtrip(self):
        import contextlib
        import io
        from dbc_compare_tool.cli import main as cli_main
        from dbc_compare_tool.core.comparator import DbcComparator
        from dbc_compare_tool.report.excel import write_excel_report
        from tests.test_signal_pairing import shifted_dbc

        old_text, new_text, _ = shifted_dbc(False)
        for side, text in (("old", old_text), ("new", new_text)):
            path = self.root / side / "bus.dbc"
            path.parent.mkdir()
            path.write_text(text, encoding="utf-8")
        # Include a parse error on each side, testing completion metadata as well.
        for side in ("old", "new"):
            (self.root / side / "broken.dbc").write_text("invalid dbc", encoding="utf-8")
        old, new = self.root / "old", self.root / "new"
        result = DbcComparator().compare_folders(old, new)
        api_path = write_excel_report(result, self.root / "api.xlsx")
        cli_path = self.root / "cli.xlsx"
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(cli_main(["--old", str(old), "--new", str(new), "--out", str(cli_path)]), 3)
        self.window.old_input.setText(str(old))
        self.window.new_input.setText(str(new))
        gui_path = self.root / "gui.xlsx"
        self.window.output_input.setText(str(gui_path))
        with patch(f"{_UI}.QMessageBox.critical") as error:
            self.window._run_compare()
            self._wait_until(lambda: self.window.run_button.isEnabled())
        error.assert_not_called()
        self.assertIn("Report incomplete", self.window.run_status.text())
        self.assertIn("2 file errors", self.window.run_status.text())
        self.assertIn(f"{result.review_count} pairs require review", self.window.run_status.text())
        def contents(path):
            wb = load_workbook(path)
            try:
                return {name: list(wb[name].values)[2 if name == "Summary" else 0:]
                        for name in wb.sheetnames}
            finally:
                wb.close()
        self.assertEqual(contents(api_path), contents(cli_path))
        self.assertEqual(contents(api_path), contents(gui_path))

    def test_removed_manual_mapping_exports_without_auto_override(self):
        old, new = self.root / "old", self.root / "new"
        old.mkdir()
        new.mkdir()
        (old / "bus.dbc").write_text(_DBC.format(sender="OldNode"), encoding="utf-8")
        (new / "bus.dbc").write_text(_DBC.format(sender="OldNode"), encoding="utf-8")
        self.window.old_input.setText(str(old))
        self.window.new_input.setText(str(new))
        self.window._saved_pair_map = {"bus.dbc": None}
        self.window._saved_pair_folders = self.window._current_folder_inputs()
        report = self.root / "manual-removed.xlsx"
        self.window.output_input.setText(str(report))
        with patch(f"{_UI}.QMessageBox.critical") as error:
            self.window._run_compare()
            self._wait_until(lambda: self.window.run_button.isEnabled())
        error.assert_not_called()
        self.assertEqual(self.window.worker.pair_map, {"bus.dbc": None})
        wb = load_workbook(report)
        try:
            rows = list(wb["Message Details"].values)
            index = rows[0].index("Change Type")
            self.assertEqual({r[index] for r in rows[1:]}, {"Added", "Removed"})
        finally:
            wb.close()


if __name__ == "__main__":
    unittest.main()
