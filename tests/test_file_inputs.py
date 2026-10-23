"""Direct DBC file comparisons and shared CLI input validation."""
from __future__ import annotations

import contextlib
import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from dbc_compare_tool.cli import main
from dbc_compare_tool.core.comparator import DbcComparator
from dbc_compare_tool.core.parser import DbcParseError


class TestFileInputs(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.old = self.root / "alpha input.dbc"
        self.new = self.root / "omega input.DBC"
        self.old.write_text('VERSION ""\n\nBO_ 100 Status: 8 OLD\n SG_ Speed : 0|8@1+ (1,0) [0|255] "" OLD\n', encoding="utf-8")
        self.new.write_text(self.old.read_text(encoding="utf-8").replace("OLD", "NEW"), encoding="utf-8")

    def _cli(self, *extra: str, old: Path | None = None, new: Path | None = None):
        stdout, stderr = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
            status = main(["--old", str(old or self.old), "--new", str(new or self.new), *extra])
        return status, stdout.getvalue(), stderr.getvalue()

    def test_explicit_files_pair_even_with_unrelated_content(self):
        self.new.write_text('VERSION ""\n\nBO_ 900 Other: 3 ECU\n', encoding="utf-8")
        progress = []
        result = DbcComparator().compare_paths(self.old, self.new, progress.append)
        self.assertEqual(len(result.file_pairs), 1)
        pair = result.file_pairs[0]
        self.assertEqual((pair.old_path, pair.new_path), (str(self.old.resolve()), str(self.new.resolve())))
        self.assertEqual(pair.status, "Manually Paired")
        self.assertIn("Explicit file selection", pair.pairing_reasons)
        self.assertEqual({c.change_type for c in result.message_changes}, {"Added", "Removed"})
        self.assertTrue(progress)

    def test_direct_file_mapping_and_include_unchanged(self):
        result = DbcComparator(include_unchanged=True, ecu_map={"OLD": "NEW"}).compare_files(self.old, self.new)
        self.assertEqual(result.summary()["Total Changes"], 0)
        self.assertEqual(result.summary()["Messages Unchanged"], 1)
        self.assertEqual(result.summary()["Signals Unchanged"], 1)

    def test_invalid_input_modes_and_missing_paths(self):
        text = self.root / "not-dbc.txt"
        text.write_text("", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "cannot be mixed"):
            DbcComparator().compare_paths(self.old, self.root)
        with self.assertRaisesRegex(ValueError, "folder or a .dbc file"):
            DbcComparator().compare_paths(text, self.new)
        with self.assertRaisesRegex(ValueError, "requires a .dbc file"):
            DbcComparator().compare_files(self.root, self.new)
        with self.assertRaisesRegex(FileNotFoundError, "does not exist"):
            DbcComparator().compare_paths(self.root / "missing.dbc", self.new)

    def test_invalid_selected_file_is_a_parse_error(self):
        self.new.write_text("not a DBC", encoding="utf-8")
        with self.assertRaises(DbcParseError):
            DbcComparator().compare_files(self.old, self.new)
        status, _, error = self._cli()
        self.assertEqual(status, 1)
        self.assertIn("Unable to parse DBC", error)

    def test_cli_files_mapping_and_default_report_name(self):
        status, output, error = self._cli("--ecu-pair", "OLD=NEW", "--include-unchanged")
        self.assertEqual((status, error), (0, ""))
        self.assertIn("Total changes: 0", output)
        report = self.root / "compared_omega input.xlsx"
        self.assertTrue(report.is_file())
        workbook = load_workbook(report)
        self.addCleanup(workbook.close)
        self.assertNotIn("Changed Properties", workbook.sheetnames)
        for name in ("Message Details", "Signal Details"):
            self.assertIn("Changed Properties", next(workbook[name].values))

    def test_cli_repeatable_mapping(self):
        self.old.write_text(self.old.read_text(encoding="utf-8").replace('"" OLD', '"" OLD_RX'), encoding="utf-8")
        self.new.write_text(self.new.read_text(encoding="utf-8").replace('"" NEW', '"" NEW_RX'), encoding="utf-8")
        status, output, error = self._cli("--ecu-pair", "OLD=NEW", "--ecu-pair", "OLD_RX=NEW_RX")
        self.assertEqual((status, error), (0, ""))
        self.assertIn("Total changes: 0", output)

    def test_cli_mapping_errors_are_actionable_and_do_not_create_reports(self):
        cases = [
            (["OLD:NEW"], "use OLD_NODE=NEW_NODE"),
            (["OLD="], "use OLD_NODE=NEW_NODE"),
            (["OLD=NEW=EXTRA"], "use OLD_NODE=NEW_NODE"),
            (["OLD=NEW", "OLD=NEW"], "OLD ECU node paired more than once"),
            (["OLD=NEW", "OTHER=NEW"], "NEW ECU nodes paired more than once"),
            (["OLD=MISSING"], "NEW: MISSING"),
        ]
        for pairs, message in cases:
            with self.subTest(pairs=pairs):
                arguments = [item for pair in pairs for item in ("--ecu-pair", pair)]
                status, _, error = self._cli(*arguments)
                self.assertEqual(status, 2)
                self.assertIn(message, error)
        self.assertFalse(list(self.root.glob("*.xlsx")))

    def test_cli_mixed_mode_is_usage_error(self):
        status, _, error = self._cli(new=self.root)
        self.assertEqual(status, 2)
        self.assertIn("cannot be mixed", error)


if __name__ == "__main__":
    unittest.main()
