"""Completeness and inventory checks retained with classic matching."""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from dbc_compare_tool.core.comparator import DbcComparator
from dbc_compare_tool.report.excel import write_excel_report

GOOD = 'VERSION ""\nBU_: ECU\nBO_ 100 Status: 8 ECU\n SG_ Value : 0|8@1+ (1,0) [0|255] "" ECU\n'


class ClassicCompletenessTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.old, self.new = self.root / "old", self.root / "new"
        self.old.mkdir()
        self.new.mkdir()

    def test_auto_fixed_path_with_unreadable_peer_is_not_added_or_removed(self):
        (self.old / "bus.dbc").write_text("broken dbc", encoding="utf-8")
        (self.new / "bus.dbc").write_text(GOOD, encoding="utf-8")
        result = DbcComparator().compare_folders(self.old, self.new)
        self.assertEqual(result.completeness, "Incomplete")
        self.assertEqual((result.diagnostics[0].side, result.diagnostics[0].path), ("OLD", "bus.dbc"))
        self.assertEqual(result.summary()["Total Changes"], 0)
        for kind in ("dbc", "message", "signal"):
            self.assertEqual(result.coverage[f"NEW {kind}"], {"not compared": 1})
            self.assertIsNone(result.coverage[f"OLD {kind}"]["unknown"])
        self.assertEqual([(p.status, p.message_count_old, p.message_count_new) for p in result.file_pairs],
                         [("Parse Error", None, 1)])
        path = write_excel_report(result, self.root / "report.xlsx")
        workbook = load_workbook(path, read_only=True)
        try:
            summary = dict(workbook["Summary"].iter_rows(min_row=4, values_only=True))
            self.assertEqual(summary["Completeness"], "Incomplete")
            self.assertIn("not compared: 1", summary["Coverage NEW message"])
            self.assertIn("Unknown", summary["Coverage OLD message"])
        finally:
            workbook.close()

    def test_manual_fixed_failure_preserves_valid_peer(self):
        (self.old / "old.dbc").write_text(GOOD, encoding="utf-8")
        (self.new / "new.dbc").write_text("broken dbc", encoding="utf-8")
        result = DbcComparator().compare_manual(self.old, self.new, {"old.dbc": "new.dbc"})
        self.assertEqual(result.summary()["Total Changes"], 0)
        self.assertEqual(result.coverage["OLD message"], {"not compared": 1})
        self.assertEqual([(p.status, p.message_count_old, p.message_count_new) for p in result.file_pairs],
                         [("Parse Error", 1, None)])

    def test_identical_files_do_not_generate_review_output(self):
        for folder in (self.old, self.new):
            (folder / "bus.dbc").write_text(GOOD, encoding="utf-8")
        result = DbcComparator().compare_files(self.old / "bus.dbc", self.new / "bus.dbc")
        self.assertEqual(result.summary()["Total Changes"], 0)
        self.assertEqual(result.coverage["OLD signal"], {"paired": 1})
        path = write_excel_report(result, self.root / "report.xlsx")
        workbook = load_workbook(path, read_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["Summary", "Message Details", "Signal Details"])
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
