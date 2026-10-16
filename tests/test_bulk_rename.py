"""Synthetic bulk-renaming regression through parsing, comparison and export."""

from __future__ import annotations

import contextlib
import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from dbc_compare_tool.cli import main
from dbc_compare_tool.core.comparator import DbcComparator
from dbc_compare_tool.report.excel import write_excel_report


# Synthetic data, not a production transport-protocol definition.
SUFFIXES = (
    "TotalMsgSize", "TotalMsgSizeEoMA", "TotalMsgSizeBAM", "TotalPktNum",
    "TotalPktNumEoMA", "TotalPktNumBAM", "PGNum_RTS", "PGNum_CTS",
    "PGNum_EoMA", "PGNum_BAM", "PGNum_Abort",
)
OLD_DBC = '''VERSION ""
BU_: ECU Receiver
BO_ 100 VCU_TP_CM: 8 ECU
 SG_ Mode M : 0|8@1+ (1,0) [0|255] "" Receiver
''' + "".join(
    f' SG_ VCU_{suffix} m{i} : 8|16@1+ (1,0) [0|65535] "" Receiver\n'
    for i, suffix in enumerate(SUFFIXES)
)
NEW_DBC = OLD_DBC.replace("VCU_", "IVC_")


class TestBulkRename(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.old_root, self.new_root = self.root / "old", self.root / "new"
        self.old = self.old_root / "nested" / "VCAN_v1.dbc"
        self.new = self.new_root / "other" / "VCAN_v2.dbc"
        for path, content in ((self.old, OLD_DBC), (self.new, NEW_DBC)):
            path.parent.mkdir(parents=True)
            path.write_text(content, encoding="utf-8")
        self.label = "VCAN_v1.dbc -> VCAN_v2.dbc"

    def _assert_renames(self, result):
        self.assertEqual([(c.change_type, c.old_name, c.new_name) for c in result.message_changes],
                         [("Renamed", "VCU_TP_CM", "IVC_TP_CM")])
        self.assertEqual({(c.old_name, c.new_name) for c in result.signal_changes},
                         {(f"VCU_{suffix}", f"IVC_{suffix}") for suffix in SUFFIXES})
        for change in result.signal_changes:
            self.assertEqual(change.change_type, "Renamed")
            self.assertEqual(change.property_diffs, ())
            self.assertEqual(change.confidence, 1.0)
            self.assertEqual(change.old_signal.multiplexer_ids, change.new_signal.multiplexer_ids)
        self.assertTrue(all(c.dbc_file == self.label for c in result.message_changes + result.signal_changes))
        self.assertEqual(result.file_pairs[0].dbc_file, self.label)

    def test_file_auto_folder_and_manual_folder_comparison(self):
        comparator = DbcComparator()
        for mode, result in (
            ("files", comparator.compare_files(self.old, self.new)),
            ("auto", comparator.compare_folders(self.old_root, self.new_root)),
            ("manual", comparator.compare_manual(self.old_root, self.new_root,
                                                 {"nested/VCAN_v1.dbc": "other/VCAN_v2.dbc"})),
        ):
            with self.subTest(mode=mode):
                self._assert_renames(result)
                self.assertIn("nested", result.file_pairs[0].old_path)
                self.assertIn("other", result.file_pairs[0].new_path)

    def test_cli_report_has_renames_and_filename_only_labels(self):
        report = self.root / "report.xlsx"
        with contextlib.redirect_stdout(io.StringIO()):
            status = main(["--old", str(self.old), "--new", str(self.new), "--out", str(report)])
        self.assertEqual(status, 0)
        workbook = load_workbook(report)
        self.addCleanup(workbook.close)
        for sheet_name in ("DBC Overview", "Message Details", "Signal Details"):
            rows = list(workbook[sheet_name].values)
            self.assertTrue(all(row[0] == self.label for row in rows[1:]))
        signals = list(workbook["Signal Details"].values)
        self.assertEqual(len(signals) - 1, len(SUFFIXES))
        self.assertTrue(all(row[2] == "Renamed" for row in signals[1:]))
        self.assertEqual(workbook["Changed Properties"].max_row, 1)

    def test_real_property_change_remains_visible_after_bulk_rename(self):
        self.new.write_text(NEW_DBC.replace('[0|65535]', '[0|60000]', 1), encoding="utf-8")
        result = DbcComparator().compare_files(self.old, self.new)
        change = next(c for c in result.signal_changes if c.new_name == "IVC_TotalMsgSize")
        self.assertEqual(change.change_type, "Renamed")
        self.assertIn(("Max", "65535", "60000"), change.property_diffs)
        report = write_excel_report(result, self.root / "changed.xlsx")
        workbook = load_workbook(report)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook["Changed Properties"].cell(2, 1).value, self.label)

    def test_same_filename_still_shows_old_and_new_without_folders(self):
        target = self.new_root / "nested" / self.old.name
        target.parent.mkdir()
        self.new.rename(target)
        result = DbcComparator().compare_folders(self.old_root, self.new_root)
        expected = "VCAN_v1.dbc -> VCAN_v1.dbc"
        self.assertEqual(result.file_pairs[0].dbc_file, expected)
        self.assertTrue(all(c.dbc_file == expected for c in result.signal_changes))

    def test_unpaired_and_parse_error_labels_have_no_folders(self):
        self.new.unlink()
        invalid = self.new.parent / "invalid.dbc"
        invalid.write_text("invalid DBC", encoding="utf-8")
        result = DbcComparator().compare_folders(self.old_root, self.new_root)
        self.assertEqual({p.dbc_file for p in result.file_pairs}, {"VCAN_v1.dbc", "invalid.dbc"})
        self.assertEqual({p.status for p in result.file_pairs}, {"DBC Removed", "Parse Error"})
        self.assertTrue(all(c.dbc_file == "VCAN_v1.dbc" for c in result.signal_changes))


if __name__ == "__main__":
    unittest.main()
