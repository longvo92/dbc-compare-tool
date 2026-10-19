"""Full inventory coverage and ECU context in impact-review reports (no Qt)."""

from __future__ import annotations

import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook

from dbc_compare_tool.cli import main
from dbc_compare_tool.core.comparator import DbcComparator, filter_result
from dbc_compare_tool.core.models import DbcDatabase, Message, Signal
from dbc_compare_tool.report.excel import write_excel_report


def _signal(name: str, start: int = 0, receivers: tuple[str, ...] = ("ECM",)) -> Signal:
    return Signal(name, start, 8, 1, "unsigned", False, 1.0, 0.0, 0, 255, "km/h", receivers)


def _database(*messages: Message) -> DbcDatabase:
    return DbcDatabase(Path("bus.dbc"), {message.name: message for message in messages})


def _rows(sheet) -> list[dict]:
    header = next(sheet.iter_rows(values_only=True))
    return [dict(zip(header, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]


_DBC = '''VERSION ""
BO_ 100 Status: 8 BCM
 SG_ Speed : 0|8@1+ (1,0) [0|255] "km/h" ECM,IC
 SG_ State : 8|8@1+ (1,0) [0|255] "" ECM
'''


class ImpactReportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.message = Message(
            "Status", 100, 8, "BCM", signals={
                "Speed": _signal("Speed", receivers=("ECM", "IC")),
                "State": _signal("State", start=8),
            }, cycle_time_ms=10, comment="Vehicle status",
        )

    def _compare(self, old=None, new=None, full=True):
        return DbcComparator(include_unchanged=full).compare_databases(
            "bus.dbc", old or _database(self.message), new or _database(self.message)
        )

    def _report(self, result):
        output = self.root / "report.xlsx"
        write_excel_report(result, output)
        workbook = load_workbook(output)
        self.addCleanup(workbook.close)
        return workbook

    def test_identical_baselines_default_empty_full_inventory_and_zero_changes(self):
        default = self._compare(full=False)
        self.assertEqual(default.message_changes, [])
        self.assertEqual(default.signal_changes, [])
        full = self._compare()
        self.assertEqual([c.change_type for c in full.message_changes], ["Unchanged"])
        self.assertEqual([c.change_type for c in full.signal_changes], ["Unchanged", "Unchanged"])
        self.assertEqual(full.summary()["Total Changes"], 0)
        self.assertEqual(full.summary()["Messages Unchanged"], 1)
        self.assertEqual(full.summary()["Signals Unchanged"], 2)
        workbook = self._report(full)
        self.assertEqual(workbook["Changed Properties"].max_row, 1)
        self.assertEqual(workbook["Message Details"]["B2"].fill.fgColor.rgb, "00F2F2F2")

    def test_signal_only_change_retains_message_and_unchanged_siblings(self):
        signals = dict(self.message.signals)
        signals["Speed"] = replace(signals["Speed"], factor=2)
        new = _database(replace(self.message, signals=signals))
        full = self._compare(new=new)
        default = self._compare(new=new, full=False)
        self.assertEqual(full.message_changes[0].change_type, "Unchanged")
        self.assertEqual({c.change_type for c in full.signal_changes}, {"Modified", "Unchanged"})
        self.assertEqual(full.summary()["Total Changes"], default.summary()["Total Changes"])
        self.assertEqual(full.summary()["Total Changes"], 1)

    def test_full_mode_cannot_drop_rows_through_change_filter(self):
        full = self._compare()
        self.assertIs(filter_result(full, {"Added"}), full)

    def test_ecu_receivers_and_identity_survive_excel_roundtrip(self):
        workbook = self._report(self._compare())
        message = _rows(workbook["Message Details"])[0]
        signals = {row["New Signal Name"]: row for row in _rows(workbook["Signal Details"])}
        self.assertEqual(message["ECU Node Tx"], "BCM")
        self.assertEqual(message["ECU Node Rx"], "ECM, IC")
        self.assertEqual(message["Old Message Name"], "Status")
        self.assertEqual(message["New Message Name"], "Status")
        self.assertEqual(message["CAN ID"], "0x64")
        self.assertEqual(signals["State"]["ECU Node Rx"], "ECM")
        self.assertEqual(signals["Speed"]["ECU Node Rx"], "ECM, IC")
        self.assertEqual(signals["Speed"]["CAN ID"], "0x64")
        for name in ("Message Details", "Signal Details", "Changed Properties"):
            self.assertEqual(workbook[name].auto_filter.ref, workbook[name].dimensions)

    def test_changed_routing_is_available_in_default_report_and_changed_properties(self):
        signals = {name: replace(signal, receivers=("VCU",)) for name, signal in self.message.signals.items()}
        new = _database(replace(self.message, transmitter="Gateway", signals=signals))
        workbook = self._report(self._compare(new=new, full=False))
        message = _rows(workbook["Message Details"])[0]
        signal = _rows(workbook["Signal Details"])[0]
        self.assertEqual(message["ECU Node Tx"], "BCM->Gateway")
        self.assertEqual(message["ECU Node Rx"], "ECM, IC->VCU")
        self.assertEqual(signal["ECU Node Rx"], "ECM, IC->VCU")
        diff = next(row for row in _rows(workbook["Changed Properties"]) if row["Property"] == "Receivers")
        self.assertIn(diff["ECU Node Rx"], ("ECM->VCU", "ECM, IC->VCU"))
        self.assertEqual(diff["ECU Node Tx"], "BCM->Gateway")

    def test_unchanged_signals_under_renamed_message_keep_both_parent_ids(self):
        new = _database(replace(self.message, name="StatusV2", can_id=400))
        result = self._compare(new=new)
        self.assertEqual(result.message_changes[0].change_type, "Renamed")
        self.assertEqual(len(result.signal_changes), 2)
        signal = _rows(self._report(result)["Signal Details"])[0]
        self.assertEqual(signal["Change Type"], "Unchanged")
        self.assertEqual(signal["Parent Message (Old)"], "Status")
        self.assertEqual(signal["Parent Message (New)"], "StatusV2")
        self.assertEqual(signal["CAN ID"], "0x64->0x190")
        message = _rows(self._report(result)["Message Details"])[0]
        self.assertEqual(message["Old Message Name"], "Status")
        self.assertEqual(message["New Message Name"], "StatusV2")
        self.assertEqual(message["CAN ID"], "0x64->0x190")
        for row in _rows(self._report(result)["Changed Properties"]):
            self.assertEqual((row["Old Name"], row["New Name"]), ("Status", "StatusV2"))

    def test_added_removed_entities_keep_only_available_routing_context(self):
        for old, new, kind, tx, can_id in (
            (_database(), _database(self.message), "Added", "BCM", "0x64"),
            (_database(self.message), _database(), "Removed", "BCM", "0x64"),
        ):
            with self.subTest(kind=kind):
                workbook = self._report(self._compare(old, new))
                for sheet in ("Message Details", "Signal Details"):
                    row = _rows(workbook[sheet])[0]
                    self.assertEqual(row["Change Type"], kind)
                    self.assertEqual(row["ECU Node Tx"], tx)
                    self.assertEqual(row["CAN ID"], can_id)
                    self.assertTrue(all("->" not in str(value) for value in row.values()))
                    entity = "Message" if sheet == "Message Details" else "Signal"
                    present, missing = ("New", "Old") if kind == "Added" else ("Old", "New")
                    self.assertTrue(row[f"{present} {entity} Name"])
                    self.assertIsNone(row[f"{missing} {entity} Name"])

    def test_signal_added_removed_inside_matched_parent(self):
        # Keep the replacement structurally unrelated to test actual add/remove routing.
        temperature = replace(_signal("Temperature", 32), length=16, factor=0.1, unit="degC")
        new_signals = {"Speed": self.message.signals["Speed"], "Temperature": temperature}
        result = self._compare(new=_database(replace(self.message, signals=new_signals)))
        rows = _rows(self._report(result)["Signal Details"])
        removed = next(row for row in rows if row["Change Type"] == "Removed")
        added = next(row for row in rows if row["Change Type"] == "Added")
        self.assertEqual(removed["Old Signal Name"], "State")
        self.assertIsNone(removed["New Signal Name"])
        self.assertIsNone(added["Old Signal Name"])
        self.assertEqual(added["New Signal Name"], "Temperature")
        self.assertEqual(removed["ECU Node Rx"], "ECM")
        self.assertEqual(added["ECU Node Rx"], "ECM")
        for row in (removed, added):
            self.assertEqual(row["Parent Message (Old)"], "Status")
            self.assertEqual(row["Parent Message (New)"], "Status")
            self.assertEqual(row["ECU Node Tx"], "BCM")

    def test_automatic_rename_keeps_ecu_context_and_full_mode(self):
        renamed = replace(self.message.signals["Speed"], name="VehicleSpeed")
        new = _database(replace(self.message, signals={"VehicleSpeed": renamed, "State": self.message.signals["State"]}))
        result = self._compare(new=new)
        self.assertTrue(result.include_unchanged)
        rows = _rows(self._report(result)["Signal Details"])
        renamed_row = next(row for row in rows if row["Change Type"] == "Renamed")
        self.assertEqual(renamed_row["Old Signal Name"], "Speed")
        self.assertEqual(renamed_row["New Signal Name"], "VehicleSpeed")
        self.assertEqual(renamed_row["ECU Node Rx"], "ECM, IC")
        self.assertEqual(result.summary()["Total Changes"], 1)

    def test_extended_and_standard_twins_are_both_in_inventory(self):
        extended = replace(self.message, name="ExtendedStatus", is_extended_frame=True)
        result = self._compare(_database(self.message, extended), _database(self.message, extended))
        self.assertEqual(len(result.message_changes), 2)
        self.assertEqual(len(result.signal_changes), 4)
        rows = _rows(self._report(result)["Message Details"])
        self.assertEqual({row["New Message Name"] for row in rows}, {"Status", "ExtendedStatus"})

    def test_changed_value_table_multiplexing_and_comments_export_only_in_changed_properties(self):
        speed = replace(self.message.signals["Speed"], is_multiplexer=True,
                        value_descriptions=((0, "Off"), (1, "On")), comment="Speed state")
        state = replace(self.message.signals["State"], multiplexer_ids=(0, 1), multiplexer_signal="Speed")
        message = replace(self.message, signals={"Speed": speed, "State": state})
        workbook = self._report(self._compare(new=_database(message)))
        rows = _rows(workbook["Changed Properties"])
        values = {(row["New Name"], row["Property"]): row["Value"] for row in rows}
        self.assertEqual(values, {
            ("Speed", "Multiplexer"): "No->Yes",
            ("Speed", "Value Descriptions"): "blank->0=Off, 1=On",
            ("Speed", "Description"): "->Speed state",
            ("State", "Multiplexer IDs"): "blank->0, 1",
            ("State", "Multiplexer Signal"): "blank->Speed",
        })
        for row in _rows(workbook["Signal Details"]):
            self.assertNotIn("Multiplexer (New)", row)
            self.assertNotIn("Value Descriptions (New)", row)
            self.assertNotIn("Signal Description (New)", row)

    def test_only_changed_message_properties_are_exported(self):
        new = replace(self.message, dlc=16, cycle_time_ms=20, comment="Updated status")
        workbook = self._report(self._compare(new=_database(new)))
        rows = _rows(workbook["Changed Properties"])
        self.assertEqual({row["Property"]: row["Value"] for row in rows}, {
            "DLC": "8->16",
            "Cycle Time": "10->20",
            "Description": "Vehicle status->Updated status",
        })
        self.assertTrue(all(row["Entity Type"] == "Message" for row in rows))

    def test_cleared_property_in_modified_message_keeps_transition(self):
        new = replace(self.message, comment="")
        workbook = self._report(self._compare(new=_database(new)))
        row = _rows(workbook["Changed Properties"])[0]
        self.assertEqual(row["Change Type"], "Modified")
        self.assertEqual(row["Property"], "Description")
        self.assertEqual(row["Value"], "Vehicle status->")

    def test_summary_records_selected_ecu_node_pairs(self):
        result = self._compare()
        result.ecu_map = {"ECM": "Powertrain", "BCM": "BodyController"}
        workbook = self._report(result)
        rows = list(workbook["Summary"].iter_rows(values_only=True))
        header = rows.index(("ECU Node", None))
        self.assertEqual(rows[header + 1:], [("BCM->BodyController", None), ("ECM->Powertrain", None)])

    def test_automatic_renamed_file_and_manual_pairing_preserve_full_inventory(self):
        old, new = self.root / "old", self.root / "new"
        old.mkdir()
        new.mkdir()
        (old / "bus.dbc").write_text(_DBC, encoding="utf-8")
        (new / "bus_v2.dbc").write_text(_DBC, encoding="utf-8")
        comparator = DbcComparator(include_unchanged=True)
        auto = comparator.compare_folders(old, new)
        manual = comparator.compare_manual(old, new, {"bus.dbc": "bus_v2.dbc"})
        for result in (auto, manual):
            self.assertEqual(len(result.message_changes), 1)
            self.assertEqual(len(result.signal_changes), 2)
            self.assertEqual(result.summary()["Total Changes"], 0)
        self.assertEqual(auto.file_pairs[0].status, "DBC Renamed")
        self.assertEqual(manual.file_pairs[0].status, "Manually Paired")

    def test_cli_full_mode_and_parse_error_coverage_are_recorded(self):
        old, new = self.root / "old", self.root / "new"
        old.mkdir()
        new.mkdir()
        for folder in (old, new):
            (folder / "bus.dbc").write_text(_DBC, encoding="utf-8")
            (folder / "bad.dbc").write_text("invalid dbc {{{", encoding="utf-8")
        output = self.root / "full.xlsx"
        self.assertEqual(main(["--old", str(old), "--new", str(new), "--out", str(output), "--include-unchanged"]), 0)
        workbook = load_workbook(output)
        self.addCleanup(workbook.close)
        summary = dict(workbook["Summary"].iter_rows(min_row=4, values_only=True))
        self.assertEqual(summary["Report Mode"], "Include Unchanged")
        self.assertEqual(summary["Total Changes"], 0)
        self.assertEqual(summary["Parse Errors (skipped DBCs)"], 1)
        self.assertEqual(workbook["Signal Details"].max_row, 3)
