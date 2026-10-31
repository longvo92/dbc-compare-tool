"""Excel report tests.

The writer is exercised end to end — written to disk, reopened with openpyxl,
and read back. Without this the report layer can crash or drop a sheet while
the whole suite stays green, which is exactly what a release build would then
ship.
"""

import tempfile
import unittest
from pathlib import Path
from unittest.mock import PropertyMock, patch

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from dbc_compare_tool.core.comparator import DbcComparator, filter_result
from dbc_compare_tool.core.models import Change, ComparisonResult, FilePairSummary
from dbc_compare_tool.report.excel import (
    ReportExportError,
    _write_signal_details,
    default_report_path,
    write_excel_report,
)

_EXAMPLES = Path(__file__).resolve().parents[1] / "examples"
_OLD_FOLDER = _EXAMPLES / "old"
_NEW_FOLDER = _EXAMPLES / "new"


def _rows(sheet) -> list[dict]:
    header = next(sheet.iter_rows(values_only=True))
    return [dict(zip(header, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]


class BaselineReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.result = DbcComparator().compare_folders(_OLD_FOLDER, _NEW_FOLDER)

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.output = Path(self._tmp.name) / "report.xlsx"
        write_excel_report(self.result, self.output)
        self.workbook = load_workbook(self.output)
        self.addCleanup(self.workbook.close)

    def test_workbook_contains_only_comparison_sheets(self):
        self.assertEqual(
            self.workbook.sheetnames,
            ["Summary", "DBC Overview", "Message Details", "Signal Details"],
        )

    def test_summary_total_matches_the_result(self):
        sheet = self.workbook["Summary"]
        totals = {row[0]: row[1] for row in sheet.iter_rows(min_row=4, values_only=True)}
        self.assertEqual(totals["Total Changes"], self.result.summary()["Total Changes"])
        self.assertEqual(totals["Messages Renamed"], self.result.summary()["Messages Renamed"])

    def test_every_signal_change_reaches_the_signal_sheet(self):
        rows = _rows(self.workbook["Signal Details"])
        self.assertEqual(len(rows), len(self.result.signal_changes))
        renamed = {(row["Old Signal Name"], row["New Signal Name"])
                   for row in rows if row["Change Type"] == "Renamed"}
        self.assertIn(("VehicleSpeed", "VehSpd"), renamed)

    def test_can_id_is_written_in_hexadecimal(self):
        rows = _rows(self.workbook["Message Details"])
        self.assertTrue(rows)
        self.assertTrue(all(str(row["CAN ID"]).startswith("0x") for row in rows if row["CAN ID"]))

    def test_changed_properties_combines_old_and_new_values(self):
        rows = _rows(self.workbook["Signal Details"])
        self.assertTrue(rows)
        factor_rows = [row for row in rows if row["New Signal Name"] == "EngineRPM"]
        self.assertTrue(factor_rows, "the EngineRPM factor change should appear")
        self.assertIn("Factor: 0.25->0.5", factor_rows[0]["Changed Properties"])

    def test_detail_columns_include_one_changed_properties_column(self):
        context = {"CAN ID", "ECU Node Tx", "ECU Node Rx"}
        shared = {
            "DBC File", "Change Type", "Confidence Score", "Confidence Level", "Match Reason", "Changed Properties",
        } | context
        expected = {
            "Message Details": shared | {"Old Message Name", "New Message Name"},
            "Signal Details": shared | {
                "Old Signal Name", "New Signal Name", "Parent Message (Old)", "Parent Message (New)",
            },
        }
        for name, headers in expected.items():
            with self.subTest(sheet=name):
                sheet = self.workbook[name]
                self.assertEqual(set(next(sheet.iter_rows(values_only=True))), headers)
                self.assertEqual(sheet.max_column, len(headers))
                self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)

    def test_rename_evidence_does_not_repeat_changed_property_values(self):
        for sheet in ("Message Details", "Signal Details"):
            rows = _rows(self.workbook[sheet])
            for row in rows:
                if row["Change Type"] == "Renamed":
                    self.assertTrue(row["Match Reason"])
                    self.assertNotIn(" -> ", row["Match Reason"])
                else:
                    self.assertIsNone(row["Match Reason"])

    def test_overview_lists_every_file_pair(self):
        rows = _rows(self.workbook["DBC Overview"])
        self.assertEqual(len(rows), len(self.result.file_pairs))
        self.assertEqual(rows[0]["Status"], "Matched")

    def test_overview_has_one_path_column_and_no_repeated_filename_column(self):
        sheet = self.workbook["DBC Overview"]
        self.assertEqual(next(sheet.values),
                         ("Status", "Path", "Pairing Confidence", "Messages", "Signals", "Pairing Reason"))
        self.assertEqual(sheet.max_column, 6)
        self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
        self.assertEqual(sheet["A2"].fill.fgColor.rgb, "00F2F2F2")

    def test_overview_combines_paths_and_counts_without_losing_zero(self):
        rows = _rows(self.workbook["DBC Overview"])
        for row, pair in zip(rows, self.result.file_pairs):
            with self.subTest(file=pair.dbc_file):
                self.assertEqual(row["Path"], pair.old_path if pair.old_path == pair.new_path
                                 else f"{pair.old_path}->{pair.new_path}")
                for label, old, new in (
                    ("Messages", pair.message_count_old, pair.message_count_new),
                    ("Signals", pair.signal_count_old, pair.signal_count_new),
                ):
                    self.assertEqual(row[label], old if old == new else f"{old}->{new}")

    def test_confidence_and_change_colors_follow_compacted_columns(self):
        for name in ("Message Details", "Signal Details"):
            sheet = self.workbook[name]
            headers = [cell.value for cell in sheet[1]]
            for row in sheet.iter_rows(min_row=2):
                if row[headers.index("Confidence Level")].value == "High":
                    self.assertEqual(row[headers.index("Confidence Level")].fill.fgColor.rgb, "00C6EFCE")
                    self.assertEqual(row[headers.index("Confidence Score")].alignment.horizontal, "center")
                properties = row[headers.index("Changed Properties")]
                self.assertTrue(properties.alignment.wrap_text)
                self.assertIsNone(sheet.row_dimensions[properties.row].height)
                if row[headers.index("Change Type")].value == "Modified":
                    self.assertEqual(properties.fill.fgColor.rgb, "00FFF2CC")

    def test_missing_parent_directory_is_created(self):
        nested = Path(self._tmp.name) / "a" / "b" / "report.xlsx"
        write_excel_report(self.result, nested)
        self.assertTrue(nested.is_file())


class OneSidedReportTests(unittest.TestCase):
    def test_added_removed_dbc_paths_and_counts_have_no_arrows_and_keep_zero(self):
        result = ComparisonResult(file_pairs=[
            FilePairSummary("added.dbc", "DBC Added", "", "new/added.dbc", None, 0, 2, 0, 0),
            FilePairSummary("removed.dbc", "DBC Removed", "old/removed.dbc", "", None, 1, 0, 0, 0),
            FilePairSummary("matched.dbc", "Matched", "same.dbc", "same.dbc", 1.0, 0, 1, 1, 0),
        ])
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "report.xlsx"
            write_excel_report(result, output)
            workbook = load_workbook(output)
            try:
                added, removed, matched = _rows(workbook["DBC Overview"])
                self.assertEqual((added["Path"], added["Messages"], added["Signals"]), ("new/added.dbc", 2, 0))
                self.assertEqual((removed["Path"], removed["Messages"], removed["Signals"]), ("old/removed.dbc", 1, 0))
                for row in (added, removed):
                    self.assertTrue(all("->" not in str(value) for value in row.values()))
                self.assertEqual((matched["Messages"], matched["Signals"]), ("0->1", "1->0"))
                self.assertEqual([workbook["DBC Overview"].cell(row, 1).fill.fgColor.rgb
                                  for row in (2, 3, 4)], ["00E2EFDA", "00FCE4D6", "00F2F2F2"])
            finally:
                workbook.close()


class ReportInputModeTests(unittest.TestCase):
    def test_file_reports_omit_overview_with_same_or_different_names_and_filters(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            old_folder, new_folder = root / "old", root / "new"
            old_folder.mkdir()
            new_folder.mkdir()
            old = old_folder / "bus.dbc"
            old.write_text('VERSION ""\nBO_ 100 Status: 8 ECU\n', encoding="utf-8")
            for new_name in ("bus.dbc", "renamed.dbc"):
                new = new_folder / new_name
                new.write_text(old.read_text(encoding="utf-8"), encoding="utf-8")
                for include_unchanged in (False, True):
                    with self.subTest(new_name=new_name, include_unchanged=include_unchanged):
                        result = DbcComparator(include_unchanged=include_unchanged).compare_paths(old, new)
                        result = filter_result(result, {"Modified"})
                        self.assertEqual(len(result.file_pairs), 1)
                        output = write_excel_report(result, root / "files.xlsx")
                        workbook = load_workbook(output)
                        try:
                            self.assertEqual(workbook.sheetnames, ["Summary", "Message Details", "Signal Details"])
                            self.assertEqual(workbook["Message Details"].max_row, 2 if include_unchanged else 1)
                        finally:
                            workbook.close()

    def test_folder_reports_keep_overview_for_one_pair_including_manual_pairing(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            old_folder, new_folder = root / "old", root / "new"
            old_folder.mkdir()
            new_folder.mkdir()
            for folder in (old_folder, new_folder):
                (folder / "bus.dbc").write_text('VERSION ""\nBO_ 100 Status: 8 ECU\n', encoding="utf-8")
            comparator = DbcComparator()
            results = (
                comparator.compare_paths(old_folder, new_folder),
                comparator.compare_manual(old_folder, new_folder, {"bus.dbc": "bus.dbc"}),
            )
            for result in results:
                with self.subTest(status=result.file_pairs[0].status):
                    self.assertEqual(len(result.file_pairs), 1)
                    output = write_excel_report(filter_result(result, {"Modified"}), root / "folders.xlsx")
                    workbook = load_workbook(output)
                    try:
                        self.assertEqual(workbook.sheetnames[:4],
                                         ["Summary", "DBC Overview", "Message Details", "Signal Details"])
                        self.assertEqual(workbook["DBC Overview"].max_row, 2)
                    finally:
                        workbook.close()


class ReportRowLimitTests(unittest.TestCase):
    @staticmethod
    def signal_changes(count):
        return [
            Change("bus.dbc", "Added", "", f"Signal{i}", None, "Entity added")
            for i in range(count)
        ]

    def test_bulk_rows_do_not_scan_existing_worksheet_dimensions(self):
        workbook = Workbook()
        self.addCleanup(workbook.close)
        with patch.object(
            Worksheet, "max_row", new_callable=PropertyMock,
            side_effect=AssertionError("Bulk row insertion must not scan existing cells"),
        ):
            _write_signal_details(workbook.active, self.signal_changes(1000))
        self.assertEqual(workbook.active.max_row, 1001)
        self.assertEqual(workbook.active.cell(1001, 6).value, "Signal999")

    def test_detail_row_limit_includes_header_and_preserves_existing_report(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "report.xlsx"
            result = ComparisonResult(signal_changes=self.signal_changes(63))
            with patch("dbc_compare_tool.report.excel.EXCEL_MAX_ROWS", 64):
                write_excel_report(result, output)
                workbook = load_workbook(output)
                try:
                    self.assertEqual(workbook["Signal Details"].max_row, 64)
                finally:
                    workbook.close()
                original = output.read_bytes()
                result.signal_changes.extend(self.signal_changes(1))
                with self.assertRaisesRegex(ReportExportError, "row limit"):
                    write_excel_report(result, output)
            self.assertEqual(output.read_bytes(), original)

    def test_summary_row_limit_is_checked_before_saving(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "report.xlsx"
            result = ComparisonResult()
            write_excel_report(result, output)
            workbook = load_workbook(output)
            try:
                summary_rows = workbook["Summary"].max_row
            finally:
                workbook.close()
            with patch("dbc_compare_tool.report.excel.EXCEL_MAX_ROWS", summary_rows):
                write_excel_report(result, output)
            original = output.read_bytes()
            with patch("dbc_compare_tool.report.excel.EXCEL_MAX_ROWS", summary_rows - 1):
                with self.assertRaisesRegex(ReportExportError, "row limit"):
                    write_excel_report(result, output)
            self.assertEqual(output.read_bytes(), original)


class DefaultReportPathTests(unittest.TestCase):
    def test_file_report_uses_stem_beside_new_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            new_file = Path(tmp) / "CAN network.v2.DBC"
            new_file.touch()
            self.assertEqual(default_report_path(new_file), Path(tmp) / "compared_CAN network.v2.xlsx")

    def test_folder_report_preserves_full_folder_name(self):
        with tempfile.TemporaryDirectory() as tmp:
            new_folder = Path(tmp) / "baseline.v2.dbc"
            new_folder.mkdir()
            self.assertEqual(default_report_path(new_folder), Path(tmp) / "compared_baseline.v2.dbc.xlsx")


if __name__ == "__main__":
    unittest.main()
