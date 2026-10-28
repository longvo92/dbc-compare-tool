from __future__ import annotations

from datetime import datetime
from pathlib import Path
import os
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from dbc_compare_tool.core.models import Change, ComparisonResult, FilePairSummary, Message
from dbc_compare_tool.report._style import BORDER, HEADER_BG, HEADER_FG


SUMMARY_ORDER = [
    "Messages Added",
    "Messages Removed",
    "Messages Modified",
    "Messages Renamed",
    "Signals Added",
    "Signals Removed",
    "Signals Modified",
    "Signals Renamed",
    "Total Changes",
]

_HEADER_BG = HEADER_BG
_HEADER_FG = HEADER_FG

# Row background color keyed by change_type
_CHANGE_ROW_FILL: dict[str, str] = {
    "Added":    "E2EFDA",   # light green
    "Removed":  "FCE4D6",   # light salmon
    "Modified": "FFF2CC",   # light yellow
    "Renamed":  "DDEBF7",   # light blue
    "Unchanged": "F2F2F2",  # light gray
}

# Confidence level cell highlight
_CONF_FILL: dict[str, str] = {
    "High":   "C6EFCE",   # green
    "Medium": "FFEB9C",   # yellow
    "Low":    "FFC7CE",   # pink-red
}

# Summary metric row colors (matches grouping)
_SUMMARY_METRIC_FILL: dict[str, str] = {
    "Messages Added":    "E8F8F5",
    "Messages Removed":  "FDEDEC",
    "Messages Modified": "FEF9E7",
    "Messages Renamed":  "EBF5FB",
    "Signals Added":     "E9F7EF",
    "Signals Removed":   "FDEDEC",
    "Signals Modified":  "FEF9E7",
    "Signals Renamed":   "EBF5FB",
    "Total Changes":     "F0F0F0",
}

# Overview sheet status row colors
_STATUS_FILL: dict[str, str] = {
    "Matched":         "F2F2F2",   # light gray
    "DBC Added":       "E2EFDA",   # light green
    "DBC Removed":     "FCE4D6",   # light salmon
    "DBC Renamed":     "DDEBF7",   # light blue
    "Manually Paired": "E4DFEC",   # light purple
    "Parse Error":     "FFC7CE",   # pink-red
}

_BORDER = BORDER

# Keep identity and routing context available for impact-review filtering.
_CONTEXT_HEADERS = ["CAN ID", "ECU Node Tx", "ECU Node Rx"]
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_CELL_LENGTH = 32_767


class ReportExportError(ValueError):
    pass


def _append(sheet, values):
    if any(isinstance(value, str) and len(value) > EXCEL_MAX_CELL_LENGTH for value in values):
        raise ReportExportError(f"Excel cell text limit exceeded in {sheet.title}")
    if sheet.max_row >= EXCEL_MAX_ROWS:
        raise ReportExportError(f"Excel row limit exceeded in {sheet.title}")
    sheet.append(values)


def _review_rows(result):
    # Include every competing edge, even if it is far below the selected score.
    selected_review = [p for p in result.pairings if p.requires_review]
    involved = {(p.old, p.new) for p in selected_review}
    involved.update((a, b) for p in selected_review for a, b, _ in p.competitors)
    pairs = [p for p in result.pairing_candidates if (p.old, p.new) in involved or not p.selected]
    return sorted(pairs + result.move_suggestions, key=lambda p: p.pair_id)


def default_report_path(new_path: Path) -> Path:
    """Place an automatically named report beside the new baseline file/folder."""
    name = new_path.stem if new_path.is_file() else new_path.name
    return new_path.parent / f"compared_{name}.xlsx"


def write_excel_report(result: ComparisonResult, output_path: Path) -> Path:
    review = _review_rows(result)
    for count in (len(result.message_changes), len(result.signal_changes), len(result.file_pairs), len(review)):
        if count + 1 > EXCEL_MAX_ROWS:
            raise ReportExportError("Excel row limit exceeded; all eligible review candidates must be exported")
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    if not result.is_file_comparison:
        overview_sheet = workbook.create_sheet("DBC Overview")
        _write_overview(overview_sheet, result.file_pairs)
        _format_overview_sheet(overview_sheet)
    message_sheet = workbook.create_sheet("Message Details")
    signal_sheet = workbook.create_sheet("Signal Details")

    _write_summary(summary_sheet, result)
    _write_message_details(message_sheet, result.message_changes)
    _write_signal_details(signal_sheet, result.signal_changes)

    _format_summary(summary_sheet)
    _format_detail_sheet(message_sheet, change_col=2, conf_score_col=6, conf_level_col=7)
    _format_detail_sheet(signal_sheet, change_col=4, conf_score_col=7, conf_level_col=8)

    if review:
        sheet = workbook.create_sheet("Pairing Review")
        _append(sheet, ["Pair ID", "Category", "Entity", "OLD Identity", "NEW Identity", "OLD DBC Path", "NEW DBC Path",
                      "OLD Parent", "NEW Parent", "Similarity", "Confidence Score", "Confidence Level",
                      "Selected", "Review Required", "Reason"])
        for pair in review:
            _append(sheet, [pair.pair_id, pair.category, (pair.old or pair.new).kind,
                pair.old.display() if pair.old else "", pair.new.display() if pair.new else "",
                pair.old.dbc_path if pair.old else "", pair.new.dbc_path if pair.new else "",
                str(pair.old.parent) if pair.old else "", str(pair.new.parent) if pair.new else "",
                pair.similarity, pair.confidence, pair.confidence_level,
                "Yes" if pair.selected else "No", "Yes" if pair.requires_review else "No", "; ".join(pair.evidence)])
        _format_detail_sheet(sheet, change_col=2, conf_score_col=11, conf_level_col=12)

    for sheet in workbook:
        if sheet.max_row > EXCEL_MAX_ROWS:
            raise ReportExportError("Excel row limit exceeded")
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    if len(cell.value) > EXCEL_MAX_CELL_LENGTH:
                        raise ReportExportError(f"Excel cell text limit exceeded in {sheet.title}")
                    cell.data_type = "s"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(suffix=".xlsx", dir=output_path.parent)
    os.close(descriptor)
    try:
        workbook.save(temporary)
        os.replace(temporary, output_path)
    finally:
        Path(temporary).unlink(missing_ok=True)
        workbook.close()
    return output_path


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------

def _write_summary(sheet, result: ComparisonResult) -> None:
    generated = datetime.now().strftime("%Y-%m-%d %H:%M")
    _append(sheet, ["DBC Compare Tool  —  Comparison Report", generated])
    _append(sheet, [])                                           # blank separator
    _append(sheet, ["Metric", "Count"])
    summary = result.summary()
    for metric in SUMMARY_ORDER:
        _append(sheet, [metric, summary.get(metric, 0)])
    for metric in summary:
        if metric not in SUMMARY_ORDER:
            _append(sheet, [metric, summary[metric]])
    _append(sheet, ["Report Mode", "Include Unchanged" if result.include_unchanged else "Changes only"])
    _append(sheet, ["Parse Errors (skipped DBCs)", len(result.diagnostics)])
    _append(sheet, ["Completeness", result.completeness])
    _append(sheet, ["Matching Policy", result.policy_version])
    _append(sheet, ["Pairs Requiring Review", result.review_count])
    _append(sheet, ["Confidence", "Heuristic; review competing candidates and parent uncertainty"])
    for group, counts in result.coverage.items():
        text = "; ".join(f"{key}: {counts.get(key, 0)}" for key in ("paired", "unmatched", "not compared"))
        if "unknown" in counts:
            text += "; unreadable entities: Unknown"
        _append(sheet, [f"Coverage {group}", text])
    for diagnostic in result.diagnostics:
        _append(sheet, [f"Parse Error {diagnostic.side} {diagnostic.path}", diagnostic.reason])
    if result.ecu_map:
        _append(sheet, [])
        _append(sheet, ["ECU Node"])
        for old_node, new_node in sorted(result.ecu_map.items()):
            _append(sheet, [_fmt_transition(old_node, new_node)])


def _write_overview(sheet, file_pairs: list[FilePairSummary]) -> None:
    _append(sheet, [
        "Status",
        "Path",
        "Pairing Confidence",
        "Messages",
        "Signals",
        "Pairing Reason",
    ])
    for fp in file_pairs:
        _append(sheet, [
            fp.status,
            _fmt_transition(fp.old_path, fp.new_path, change_type=fp.status),
            _fmt_confidence(fp.pairing_confidence),
            _fmt_transition(fp.message_count_old if fp.message_count_old is not None else "Unknown",
                            fp.message_count_new if fp.message_count_new is not None else "Unknown", change_type=fp.status),
            _fmt_transition(fp.signal_count_old if fp.signal_count_old is not None else "Unknown",
                            fp.signal_count_new if fp.signal_count_new is not None else "Unknown", change_type=fp.status),
            "; ".join(fp.pairing_reasons),
        ])


def _write_message_details(sheet, changes: list[Change]) -> None:
    _append(sheet, [
        "DBC File",
        "Change Type",
        "Old Message Name",
        "New Message Name",
        "CAN ID",
        "Confidence Score",
        "Confidence Level",
        "Match Reason",
        *_CONTEXT_HEADERS[1:],
        "Changed Properties",
        "Pair ID", "Review Required",
    ])
    for change in changes:
        can_id, transmitter, receivers = _context_values(change)
        _append(sheet, [
            _dbc_label(change.dbc_file),
            change.change_type,
            change.old_name,
            change.new_name,
            can_id or _fmt_can_id(change.can_id),
            _fmt_confidence(change.confidence),
            change.confidence_level,
            _match_reason(change),
            transmitter,
            receivers,
            _changed_properties(change),
            change.pairing.pair_id if change.pairing else "",
            "Yes" if change.pairing and change.pairing.requires_review else "No",
        ])


def _write_signal_details(sheet, changes: list[Change]) -> None:
    _append(sheet, [
        "DBC File",
        "Parent Message (Old)",
        "Parent Message (New)",
        "Change Type",
        "Old Signal Name",
        "New Signal Name",
        "Confidence Score",
        "Confidence Level",
        "Match Reason",
        *_CONTEXT_HEADERS,
        "Changed Properties",
        "Pair ID", "Review Required",
    ])
    for change in changes:
        _append(sheet, [
            _dbc_label(change.dbc_file),
            *_parent_message_names(change),
            change.change_type,
            change.old_name,
            change.new_name,
            _fmt_confidence(change.confidence),
            change.confidence_level,
            _match_reason(change),
            *_context_values(change, signal=True),
            _changed_properties(change),
            change.pairing.pair_id if change.pairing else "",
            "Yes" if change.pairing and change.pairing.requires_review else "No",
        ])


# ---------------------------------------------------------------------------
# Formatters
# ---------------------------------------------------------------------------

def _format_summary(sheet) -> None:
    # Row 1 — report title
    title_cell = sheet.cell(row=1, column=1)
    title_cell.font = Font(bold=True, size=13, color="1F4E78")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 26

    date_cell = sheet.cell(row=1, column=2)
    date_cell.font = Font(italic=True, size=10, color="595959")
    date_cell.alignment = Alignment(horizontal="right", vertical="center")

    # Row 3 — column headers
    header_fill = PatternFill("solid", fgColor=_HEADER_BG)
    header_font = Font(color=_HEADER_FG, bold=True, size=11)
    for cell in sheet[3]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER
    sheet.row_dimensions[3].height = 22

    # Rows 4+ — metric data
    for row_idx in range(4, sheet.max_row + 1):
        metric = sheet.cell(row_idx, 1).value
        row = sheet[row_idx]
        bg = _SUMMARY_METRIC_FILL.get(metric, "FFFFFF")
        fill = PatternFill("solid", fgColor=bg)
        for cell in row:
            cell.fill = fill
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="center")
        if metric == "Total Changes":
            for cell in row:
                cell.font = Font(bold=True, size=11)
        elif metric == "ECU Node":
            for cell in row:
                cell.fill = header_fill
                cell.font = header_font
        if metric == "Completeness" and sheet.cell(row_idx, 2).value == "Incomplete":
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="FFC7CE")
                cell.font = Font(bold=True, color="9C0006")
        sheet.row_dimensions[row_idx].height = 18

    for column, default_width in (("A", 30), ("B", 24)):
        max_len = max(len(str(cell.value or "")) for cell in sheet[column][2:])
        sheet.column_dimensions[column].width = min(max(default_width, max_len + 2), 60)
    sheet.freeze_panes = "A4"


def _format_overview_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor=_HEADER_BG)
    header_font = Font(color=_HEADER_FG, bold=True, size=11)

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
    sheet.row_dimensions[1].height = 26

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        status = str(row[0].value or "")
        row_fill = PatternFill("solid", fgColor=_STATUS_FILL.get(status, "FFFFFF"))
        for cell in row:
            cell.fill = row_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = _BORDER
        sheet.row_dimensions[row_idx].height = 18

    for col_cells in sheet.columns:
        max_len = max(len(str(cell.value or "")) for cell in col_cells)
        col_letter = get_column_letter(col_cells[0].column)
        sheet.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 60)


def _format_detail_sheet(
    sheet,
    change_col: int,
    conf_score_col: int,
    conf_level_col: int,
) -> None:
    header_fill = PatternFill("solid", fgColor=_HEADER_BG)
    header_font = Font(color=_HEADER_FG, bold=True, size=11)

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
    sheet.row_dimensions[1].height = 26

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for row in sheet.iter_rows(min_row=2):
        change_type = str(row[change_col - 1].value or "")
        row_bg = _CHANGE_ROW_FILL.get(change_type, "FFFFFF")
        row_fill = PatternFill("solid", fgColor=row_bg)

        for cell in row:
            cell.fill = row_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = _BORDER

        # Highlight confidence level cell individually
        conf_cell = row[conf_level_col - 1]
        conf_level = str(conf_cell.value or "")
        if conf_level in _CONF_FILL:
            conf_cell.fill = PatternFill("solid", fgColor=_CONF_FILL[conf_level])
            conf_cell.font = Font(bold=True)
            conf_cell.alignment = Alignment(horizontal="center", vertical="top")

        # Center-align confidence score
        row[conf_score_col - 1].alignment = Alignment(horizontal="center", vertical="top")

    # Auto-width columns (capped at 60)
    for col_cells in sheet.columns:
        max_len = max(len(str(cell.value or "")) for cell in col_cells)
        col_letter = get_column_letter(col_cells[0].column)
        sheet.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 60)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _changed_properties(change: Change) -> str:
    """Keep all property deltas in the entity's detail row."""
    return "; ".join(
        f"{prop}: {_fmt_transition(old, new, change_type=change.change_type)}"
        for prop, old, new in change.property_diffs
    )


def _fmt_transition(
    old: str | int | None, new: str | int | None, *, change_type: str = "",
) -> str | int:
    """Show only the existing side for additions/removals; retain other deltas."""
    old = "" if old is None else old
    new = "" if new is None else new
    if change_type in {"Added", "DBC Added"}:
        return new
    if change_type in {"Removed", "DBC Removed"}:
        return old
    return old if old == new else f"{old}->{new}"


def _dbc_label(label: str) -> str:
    old, separator, new = label.partition(" -> ")
    return _fmt_transition(old, new) if separator else label


def _parent_message_names(change: Change) -> tuple[str, str]:
    if change.old_message is None and change.new_message is None:
        return (
            change.parent_message if change.change_type != "Added" else "",
            change.parent_message if change.change_type != "Removed" else "",
        )
    return (
        change.old_message.name if change.old_message else "",
        change.new_message.name if change.new_message else "",
    )


def _match_reason(change: Change) -> str:
    """Retain rename evidence without copying technical property changes."""
    if change.change_type != "Renamed":
        return ""
    return "; ".join(change.pairing.evidence) if change.pairing else ""


def _message_receivers(message: Message | None) -> str:
    if message is None:
        return ""
    return ", ".join(sorted({node for signal in message.signals.values() for node in signal.receivers}))


def _context_values(change: Change, signal: bool = False) -> list[str]:
    old, new = change.old_message, change.new_message
    if signal:
        old_rx = ", ".join(change.old_signal.receivers) if change.old_signal else ""
        new_rx = ", ".join(change.new_signal.receivers) if change.new_signal else ""
    else:
        old_rx, new_rx = _message_receivers(old), _message_receivers(new)
    return [
        _fmt_transition(
            _fmt_can_id(old.can_id) if old else "", _fmt_can_id(new.can_id) if new else "",
            change_type=change.change_type,
        ),
        _fmt_transition(old.transmitter if old else "", new.transmitter if new else "", change_type=change.change_type),
        _fmt_transition(old_rx, new_rx, change_type=change.change_type),
    ]


def _fmt_confidence(confidence: float | None) -> str:
    return "" if confidence is None else f"{confidence:.2f}"


def _fmt_can_id(can_id: int | None) -> str:
    if can_id is None:
        return ""
    try:
        return f"0x{int(can_id):X}"
    except (TypeError, ValueError):
        return str(can_id)
